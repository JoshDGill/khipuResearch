"""Steward step 3: load the KFG Excel files into interim tables (plan §4.1).

Values are kept as recorded (strings), with provenance pointing at file, sheet and 1-based
sheet row. Writes data/interim/kfg/{khipus,groups,cords}.parquet.

    python -m steward.kfg
"""

from __future__ import annotations

import json
import re
import sys
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import pandas as pd

from khipu import paths

EXCEL_DIR = paths.RAW / "kfg" / "excel"
OUT_DIR = paths.INTERIM / "kfg"
GROUP_RANGE = re.compile(r"\((\d+)(?:-(\d+))?\)")
CORD_NAME = re.compile(r"^p(\d+)((?:s\d+)*)$")


def _kv(rows: list[tuple]) -> dict[str, str]:
    out = {}
    for r in rows:
        cell = r[0] if r else None
        if cell is None or str(cell).startswith("!"):
            continue
        key, _, val = str(cell).partition(":")
        out[key.strip()] = val.strip()
    return out


def parse_group_line(line: str) -> tuple[int, int] | None:
    m = GROUP_RANGE.search(line)
    if not m:
        return None
    a = int(m.group(1))
    return a, int(m.group(2) or a)


def parse_cord_name(name: str) -> tuple[int, int, str] | None:
    """'p12s3s1' -> (pendant 12, level 3, parent name 'p12s3')."""
    m = CORD_NAME.match(name)
    if not m:
        return None
    subs = re.findall(r"s\d+", m.group(2))
    parent = "p" + m.group(1) + "".join(subs[:-1]) if subs else ""
    return int(m.group(1)), 1 + len(subs), parent


def parse_value(v) -> int | None:
    if v is None:
        return None
    s = str(v).strip()
    return int(s) if re.fullmatch(r"-?\d+", s) else None


def load_file(path: Path) -> tuple[dict, list[dict], list[dict]]:
    rel = f"excel/{path.name}"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    meta = _kv(list(wb["Khipu"].iter_rows(values_only=True)))
    kh = meta.get("Name") or path.stem
    khipu = {"kh": kh, "file": rel, **{k.lower().replace(" ", "_").replace("/", "_"): v
                                       for k, v in meta.items()}}

    groups = []
    for i, row in enumerate(wb["CordGroups"].iter_rows(values_only=True), start=1):
        line = row[0] if row else None
        if line is None or str(line).startswith("!"):
            continue
        rng = parse_group_line(str(line))
        groups.append({"kh": kh, "group_index": len(groups), "first_pendant": rng and rng[0],
                       "last_pendant": rng and rng[1], "line": str(line),
                       "src_source": "kfg", "src_file": rel, "src_table": "CordGroups",
                       "src_row": str(i)})

    rows = list(wb["Cords"].iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    cords = []
    for i, r in enumerate(rows[1:], start=2):
        d = dict(zip(header, r, strict=False))
        name = str(d.get("Cord_Name") or "").strip()
        parsed = parse_cord_name(name)
        cords.append({
            "kh": kh, "cord_name": name, "row_order": i - 2,
            "pendant_num": parsed and parsed[0], "level": parsed and parsed[1],
            "parent_name": parsed and parsed[2],
            **{k.lower(): (None if d.get(k) is None else str(d.get(k))) for k in header},
            "value_int": parse_value(d.get("Value")),
            "src_source": "kfg", "src_file": rel, "src_table": "Cords", "src_row": str(i),
        })
    wb.close()
    return khipu, groups, cords


def main() -> int:
    files = sorted(EXCEL_DIR.glob("*.xlsx"))
    khipus, groups, cords, failed = [], [], [], []
    for f in files:
        try:
            k, g, c = load_file(f)
        except Exception as e:  # recorded, not silently skipped
            failed.append({"file": f.name, "error": repr(e)})
            continue
        khipus.append(k)
        groups += g
        cords += c
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    kdf = pd.DataFrame(khipus).astype(str)
    kdf.to_parquet(OUT_DIR / "khipus.parquet", index=False)
    pd.DataFrame(groups).to_parquet(OUT_DIR / "groups.parquet", index=False)
    cdf = pd.DataFrame(cords)
    cdf.to_parquet(OUT_DIR / "cords.parquet", index=False)
    pd.DataFrame(failed, columns=["file", "error"]).to_csv(OUT_DIR / "failed.csv", index=False)
    log = {"step": "steward.kfg", "at": datetime.now(UTC).isoformat(timespec="seconds"),
           "files": len(files), "khipus": len(khipus), "failed": len(failed),
           "cords": len(cdf), "unparsed_cord_names": int(cdf["pendant_num"].isna().sum()),
           "non_integer_values": int(cdf["value_int"].isna().sum())}
    with open(paths.INTERIM / "transform_log.jsonl", "a") as fh:
        fh.write(json.dumps(log) + "\n")
    print(json.dumps(log))
    return 0


if __name__ == "__main__":
    sys.exit(main())
